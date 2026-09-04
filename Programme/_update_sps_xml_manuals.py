from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCS = Path(r"C:\Neuronales Netz\Dokumente")
DE = DOCS / "NeuronNetz_Hilfe_DE_finale_Version.docx"
EN = DOCS / "NeuronNetz_Help_EN_final_Version.docx"


def paragraph_with_prefix(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"Absatz nicht gefunden: {prefix}")


def replace_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def add_marker_before(reference, text):
    paragraph_element = OxmlElement("w:p")
    reference._p.addprevious(paragraph_element)
    paragraph = reference._parent.add_paragraph()
    paragraph._p.getparent().remove(paragraph._p)
    paragraph_element.getparent().replace(paragraph_element, paragraph._p)

    paragraph.style = "Normal"
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(0x9C, 0x57, 0x00)

    ppr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "FFF2CC")
    ppr.append(shading)
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), "E6A700")
        borders.append(border)
    ppr.append(borders)
    return paragraph


def first_drawing_after(document, paragraph):
    found = False
    for candidate in document.paragraphs:
        if candidate._p is paragraph._p:
            found = True
            continue
        if found and candidate._p.xpath(".//w:drawing"):
            return candidate
    raise RuntimeError("Nachfolgende Abbildung nicht gefunden")


def update_german():
    document = Document(DE)
    neutral = paragraph_with_prefix(document, "Der Menüpunkt IEC 61131-10 XML")
    replace_text(
        neutral,
        "Der Menüpunkt IEC 61131-10 XML verwendet einen eigenen "
        "herstellerneutralen IEC-Generator. Das erzeugte XML enthält weder "
        "Mitsubishi-Bezeichnungen noch GX-Works-spezifische Funktionsaufrufe "
        "und verwendet durchgehend den Einzelvariablenaufbau. Variablenhinweise "
        "werden standardkonform als Documentation-Elemente abgelegt. Der Import "
        "in GX Works3 wurde praktisch geprüft; GX Works3 übernimmt diese neutralen "
        "Hinweise jedoch nicht in seine sichtbare Kommentarspalte.",
    )
    neutral_window = paragraph_with_prefix(document, "Das XML-Fenster besitzt")
    replace_text(
        neutral_window,
        "Das XML-Fenster besitzt einheitlich gerahmte Bereiche für "
        "Bausteineinstellungen, FB-Vorschau, technische Angaben, "
        "Sicherheitshinweis und den vollständigen schreibgeschützten XML-Code. "
        "XML kopieren übernimmt das gesamte Dokument in die Zwischenablage; "
        "XML-Datei speichern... speichert genau den angezeigten Inhalt. Wie "
        "vollständig andere Entwicklungsumgebungen IEC 61131-10 importieren, "
        "muss im jeweiligen Zielsystem praktisch geprüft werden.",
    )
    neutral_image = first_drawing_after(document, neutral_window)
    add_marker_before(
        neutral_image,
        "SCREENSHOT AKTUALISIEREN: SPS-Export-Menü und XML-Exportfenster. Im "
        "Menü muss die Reihenfolge IEC 61131-10 XML, Trennlinie, Mitsubishi GX "
        "Works3 – XML, Mitsubishi GX Works3 – ASC und Mitsubishi GX Works2 – "
        "ASC sichtbar sein.",
    )

    transfer_heading = paragraph_with_prefix(
        document, "Übertragung nach GX Works2 und GX Works3"
    )
    replace_text(transfer_heading, "Mitsubishi GX Works3 XML und ASC-Übertragung")
    transfer = paragraph_with_prefix(document, "Für GX Works2 werden")
    replace_text(
        transfer,
        "Der Menüpunkt Mitsubishi GX Works3 – XML erzeugt eine gesonderte, in "
        "GX Works3 praktisch importierte XML-Variante. Sie behält den "
        "Einzelvariablenaufbau und den GX-Works3-kompatiblen Structured Text bei, "
        "speichert Variablenkommentare aber als allgemeinen Kommentar Nr. 1 in "
        "AddData/VariableComments. Dadurch erscheinen die Kommentare nach dem "
        "Import direkt in der Kommentarspalte.",
    )
    transfer2 = paragraph_with_prefix(document, "Übertragen Sie zuerst")
    replace_text(
        transfer2,
        "Beim ASC-Transfer für GX Works3 muss im Local-Label-Editor Show Details "
        "aktiviert sein; die Reihenfolge lautet Label Name, Data Type, Class, "
        "Initial Value, Constant und Comment. Für GX Works2 lautet sie Class, "
        "Label Name, Data Type, Constant und Comment. Übertragen Sie zuerst die "
        "Deklarationen und anschließend den Structured Text. Alternativ speichert "
        "Vollständige ASC-Datei speichern... den gesamten Funktionsbaustein.",
    )
    labels_image = first_drawing_after(document, transfer2)
    add_marker_before(
        labels_image,
        "SCREENSHOT AKTUALISIEREN: GX Works3 Local Label Settings nach dem Import "
        "über Mitsubishi GX Works3 – XML. Mehrere Konstanten und die sichtbaren "
        "Variablenkommentare müssen erkennbar sein.",
    )
    document.save(DE)


def update_english():
    document = Document(EN)
    neutral = paragraph_with_prefix(document, "The IEC 61131-10 XML menu item")
    replace_text(
        neutral,
        "The IEC 61131-10 XML menu item uses a dedicated manufacturer-neutral "
        "IEC generator. The generated XML contains neither Mitsubishi "
        "designations nor GX Works-specific function calls and consistently uses "
        "the individual-variable layout. Variable notes are stored in standard "
        "Documentation elements. Import into GX Works3 has been tested in "
        "practice, but GX Works3 does not transfer these neutral notes to its "
        "visible Comment column.",
    )
    neutral_window = paragraph_with_prefix(document, "The XML window has")
    replace_text(
        neutral_window,
        "The XML window has consistently framed areas for block settings, FB "
        "preview, technical information, safety notice, and the complete "
        "read-only XML code. Copy XML places the complete document on the "
        "clipboard; Save XML file... stores exactly the displayed content. The "
        "completeness of IEC 61131-10 import in other development environments "
        "must be verified in the respective target system.",
    )
    marker = paragraph_with_prefix(document, "[INSERT SCREENSHOT: IEC 61131-10 XML")
    replace_text(
        marker,
        "UPDATE SCREENSHOT: PLC Export menu and XML export window. The menu must "
        "show IEC 61131-10 XML, a separator, Mitsubishi GX Works3 – XML, "
        "Mitsubishi GX Works3 – ASC, and Mitsubishi GX Works2 – ASC in this order.",
    )
    marker.runs[0].bold = True
    marker.runs[0].font.color.rgb = RGBColor(0x9C, 0x57, 0x00)

    transfer_heading = paragraph_with_prefix(document, "Transfer to GX Works2")
    replace_text(transfer_heading, "Mitsubishi GX Works3 XML and ASC transfer")
    transfer = paragraph_with_prefix(document, "For GX Works2, paste declarations")
    replace_text(
        transfer,
        "Mitsubishi GX Works3 – XML creates a separate XML variant that has been "
        "imported successfully into GX Works3 in practice. It retains the "
        "individual-variable layout and GX Works3-compatible Structured Text, "
        "but stores variable comments as general comment No. 1 in "
        "AddData/VariableComments. The comments therefore appear directly in the "
        "Comment column after import.",
    )
    transfer2 = paragraph_with_prefix(document, "First transfer the declarations")
    replace_text(
        transfer2,
        "For GX Works3 ASC transfer, enable Show Details in the Local Label "
        "editor; the order is Label Name, Data Type, Class, Initial Value, "
        "Constant, and Comment. For GX Works2, the order is Class, Label Name, "
        "Data Type, Constant, and Comment. Transfer declarations first and then "
        "Structured Text. Alternatively, Save complete ASC file... stores the "
        "complete function block.",
    )
    label_marker = paragraph_with_prefix(
        document, "[INSERT SCREENSHOT: GX Works3 Local Label Settings"
    )
    replace_text(
        label_marker,
        "UPDATE SCREENSHOT: GX Works3 Local Label Settings after import with "
        "Mitsubishi GX Works3 – XML. Show several constants and the visible "
        "variable comments.",
    )
    label_marker.runs[0].bold = True
    label_marker.runs[0].font.color.rgb = RGBColor(0x9C, 0x57, 0x00)
    document.save(EN)


if __name__ == "__main__":
    update_german()
    update_english()
    print(DE)
    print(EN)
