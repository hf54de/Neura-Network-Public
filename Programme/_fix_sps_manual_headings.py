from pathlib import Path
from docx import Document


FILES = (
    (
        Path(r"C:\Neuronales Netz\Dokumente\NeuronNetz_Hilfe_DE_finale_Version.docx"),
        "Übertragung nach GX Works2 und GX Works3",
        "Mitsubishi GX Works3 XML und ASC-Übertragung",
    ),
    (
        Path(r"C:\Neuronales Netz\Dokumente\NeuronNetz_Help_EN_final_Version.docx"),
        "Transfer to GX Works2 and GX Works3",
        "Mitsubishi GX Works3 XML and ASC transfer",
    ),
)


for path, old, new in FILES:
    document = Document(path)
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == old]
    if len(matches) != 1:
        raise RuntimeError(f"Erwartete genau eine verbleibende Überschrift in {path}: {len(matches)}")
    matches[0].clear()
    matches[0].add_run(new)
    document.save(path)
    print(path)
